# -*- coding: utf-8 -*-
"""Motor reutilizable para generar documentos DOCX branded de EDUVIAGT."""
import io
import os
import re

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- Paleta ----------
INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1F, 0x3B, 0x73)
GREY = RGBColor(0x55, 0x55, 0x55)
CODE_BG = "F2F2F2"
HDR_BG = "1F3B73"

LOGO_UMG = None
LOGO_EDUVI = None

def set_logos(umg_path, eduvi_path):
    global LOGO_UMG, LOGO_EDUVI
    LOGO_UMG = umg_path if os.path.exists(umg_path) else None
    LOGO_EDUVI = eduvi_path if os.path.exists(eduvi_path) else None


def _style_base(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.font.color.rgb = INK
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.08
    # east asian font
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Calibri")

    for name, size, color, before, after in (
        ("Heading 1", 16, ACCENT, 16, 6),
        ("Heading 2", 13, ACCENT, 12, 4),
        ("Heading 3", 11.5, ACCENT, 10, 3),
    ):
        s = doc.styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.color.rgb = color
        s.font.bold = True
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    # Code style
    try:
        cs = doc.styles.add_style("Code", 1)  # 1 = paragraph
    except ValueError:
        cs = doc.styles["Code"]
    cs.font.name = "Consolas"
    cs.font.size = Pt(9.5)
    cs.font.color.rgb = INK
    cs.paragraph_format.space_after = Pt(4)
    cs.paragraph_format.space_before = Pt(4)
    cs.paragraph_format.left_indent = Cm(0.6)
    cs.paragraph_format.right_indent = Cm(0.4)
    cs.paragraph_format.line_spacing = 1.0
    rpr = cs.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Consolas")
    rfonts.set(qn("w:hAnsi"), "Consolas")
    rfonts.set(qn("w:cs"), "Consolas")


def new_doc():
    doc = Document()
    _style_base(doc)
    # margins
    for s in doc.sections:
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)
    return doc


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def add_cover(doc, title, subtitle, meta_lines, version="1.0", date="2026-10-29"):
    # title page
    sec = doc.sections[0]
    # remove header/footer on title page by marking different first page
    sec.different_first_page_header_footer = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    if LOGO_UMG:
        r = p.add_run()
        r.add_picture(LOGO_UMG, width=Cm(3.2))
    if LOGO_EDUVI:
        r = p.add_run()
        r.add_picture(LOGO_EDUVI, width=Cm(2.6))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("UNIVERSIDAD MARIANO GÁLVEZ")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FACULTAD DE INGENIERÍA EN SISTEMAS DE INFORMACIÓN Y CIENCIAS DE LA COMPUTACIÓN")
    r.font.size = Pt(9.5); r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True; r.font.size = Pt(19); r.font.color.rgb = ACCENT
    p.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.italic = True; r.font.size = Pt(11.5); r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(24)

    # meta table
    if meta_lines:
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in meta_lines:
            row = t.add_row().cells
            row[0].text = ""
            row[1].text = ""
            r = row[0].paragraphs[0].add_run(k)
            r.bold = True; r.font.size = Pt(10); r.font.color.rgb = ACCENT
            r2 = row[1].paragraphs[0].add_run(v)
            r2.font.size = Pt(10)
            _shade(row[0], "EAEFF7")
        # widths
        for row in t.rows:
            row.cells[0].width = Cm(5.5)
            row.cells[1].width = Cm(11.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Versión %s · %s" % (version, date))
    r.font.size = Pt(9); r.font.color.rgb = GREY
    p.paragraph_format.space_before = Pt(18)

    doc.add_page_break()


def add_header_footer(doc, doc_name):
    for section in doc.sections:
        # header
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if LOGO_EDUVI:
            r = hp.add_run()
            r.add_picture(LOGO_EDUVI, width=Cm(1.6))
        r = hp.add_run("   " + doc_name)
        r.font.size = Pt(8.5); r.font.color.rgb = GREY
        # footer
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("EDUVIAGT · %s · Página " % doc_name)
        r.font.size = Pt(8.5); r.font.color.rgb = GREY
        # page number field
        fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
        instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve"); instrText.text = "PAGE"
        fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
        run = fp.add_run()
        run.font.size = Pt(8.5); run.font.color.rgb = GREY
        run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


def add_toc(doc, title="Tabla de contenidos"):
    doc.add_heading(title, level=1)
    p = doc.add_paragraph()
    r = p.add_run()
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve")
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "separate")
    fldChar3 = OxmlElement("w:fldChar"); fldChar3.set(qn("w:fldCharType"), "end")
    r._r.append(fldChar1); r._r.append(instrText); r._r.append(fldChar2); r._r.append(fldChar3)
    doc.add_paragraph().add_run(" (Actualice el campo en Word: clic derecho → Actualizar campo) ").italic = True
    doc.add_page_break()


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "1F3B73")
    pbdr.append(bottom); pPr.append(pbdr)
    return p


def add_code(doc, code):
    for line in code.split("\n"):
        p = doc.add_paragraph(style="Code")
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line.strip() != "" else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        # shade
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), CODE_BG)
        pPr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(str(h))
        r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        _shade(hdr[i], HDR_BG)
    ncols = len(headers)
    for row in rows:
        # normalize column count
        row = (list(row) + [""] * ncols)[:ncols]
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(9.5)
            if i == 0:
                r.bold = True
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_image(doc, path, caption=None, width=14.0):
    if not os.path.exists(path):
        p = doc.add_paragraph()
        r = p.add_run("[IMAGEN NO ENCONTRADA: %s]" % path)
        r.italic = True; r.font.color.rgb = RGBColor(0xCC,0x00,0x00)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(path, width=Cm(width))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(caption)
        cr.italic = True; cr.font.size = Pt(9); cr.font.color.rgb = GREY


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    _add_rich(p, text)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    _add_rich(p, text)
    return p


def _add_rich(p, text):
    # bold **, italic *, code `
    # split by code first
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x8B,0x00,0x00)
            continue
        # bold
        for chunk in re.split(r"(\*\*[^*]+\*\*)", part):
            if chunk.startswith("**") and chunk.endswith("**"):
                r = p.add_run(chunk[2:-2]); r.bold = True
            else:
                # italic
                for sub in re.split(r"(\*[^*]+\*)", chunk):
                    if sub.startswith("*") and sub.endswith("*") and len(sub) > 2:
                        r = p.add_run(sub[1:-1]); r.italic = True
                    else:
                        p.add_run(sub)


def md_to_docx(md_text, doc):
    """Convierte Markdown a elementos del documento."""
    lines = md_text.split("\n")
    i = 0
    in_code = False
    code_buf = []
    table_buf = []
    in_table = False

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        # parse header + separator + rows
        rows = [r.strip() for r in table_buf if r.strip()]
        if not rows:
            table_buf = []
            return
        # header
        header = [c.strip() for c in rows[0].strip().strip("|").split("|")]
        # skip separator row
        data_rows = []
        for r in rows[1:]:
            if re.match(r"^\|[\s\-:|]+\|$", r.strip()):
                continue
            data_rows.append([c.strip() for c in r.strip().strip("|").split("|")])
        # remove empty trailing
        if not data_rows:
            table_buf = []
            return
        # normalize column counts
        maxcols = max([len(header)] + [len(r) for r in data_rows])
        header = (header + [""] * maxcols)[:maxcols]
        data_rows = [(r + [""] * maxcols)[:maxcols] for r in data_rows]
        add_table(doc, header, data_rows)
        table_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # code fence
        if stripped.startswith("```"):
            if in_code:
                add_code(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # table
        if stripped.startswith("|"):
            if not in_table:
                in_table = True
            table_buf.append(line)
            i += 1
            continue
        else:
            if in_table:
                flush_table()
                in_table = False

        # blank line
        if stripped == "":
            i += 1
            continue

        # horizontal rule
        if re.match(r"^-{3,}$", stripped) or re.match(r"^\*{3,}$", stripped):
            add_hr(doc)
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 3)
            text = m.group(2).strip()
            # remove trailing #
            text = re.sub(r"\s+#*$", "", text)
            doc.add_heading(text, level=level)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            r = p.add_run(" ".join(quote_lines))
            r.italic = True; r.font.color.rgb = GREY
            p.paragraph_format.left_indent = Cm(1.0)
            continue

        # image
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", stripped)
        if m:
            alt, src = m.group(1), m.group(2)
            # resolve relative to repo root if needed
            cand = src
            if not os.path.isabs(cand):
                # try relative to docs/entregables
                base = "C:/dev/Sigeb-v1.1/docs/entregables"
                cand2 = os.path.join(base, src)
                if os.path.exists(cand2):
                    cand = cand2
                else:
                    cand2 = os.path.join("C:/dev/Sigeb-v1.1", src)
                    if os.path.exists(cand2):
                        cand = cand2
            add_image(doc, cand, caption=alt if alt else None, width=14.0)
            i += 1
            continue

        # unordered list
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            level = 0 if indent < 2 else 1
            add_bullet(doc, m.group(2).strip(), level=level)
            i += 1
            continue

        # ordered list
        m = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if m:
            add_numbered(doc, m.group(2).strip())
            i += 1
            continue

        # normal paragraph (collect consecutive non-special lines)
        para_lines = [stripped]
        j = i + 1
        while j < len(lines):
            nj = lines[j].strip()
            if (nj == "" or nj.startswith("#") or nj.startswith("|") or nj.startswith("```")
                    or nj.startswith(">") or nj.startswith("- ") or nj.startswith("* ")
                    or re.match(r"^\d+\.\s", nj) or re.match(r"^-{3,}$", nj)
                    or re.match(r"^!\[", nj)):
                break
            para_lines.append(nj)
            j += 1
        p = doc.add_paragraph()
        _add_rich(p, " ".join(para_lines))
        i = j
        continue

    if in_table:
        flush_table()
    if in_code:
        add_code(doc, "\n".join(code_buf))


def add_company_identity(doc):
    """Seccion branded: Quienes somos / Como trabajamos."""
    doc.add_heading("Quiénes somos", level=1)
    p = doc.add_paragraph()
    _add_rich(p, "EDUVIAGT es la empresa consultora seleccionada por el Ministerio de Educación de Guatemala "
                 "para analizar, diseñar e implementar la **Plataforma Nacional para la Gestión Integral de Becas**. "
                 "Somos un equipo de 6 integrantes con roles Scrum definidos, que aplica prácticas reales de ingeniería "
                 "de software: SOLID, patrones de diseño, arquitectura modular, pruebas automatizadas y control de "
                 "versiones con Git.")
    doc.add_heading("Cómo trabajamos", level=2)
    add_bullet(doc, "**Metodología:** Scrum con sprints de 2 semanas, 9 sprints en total (agosto 2026 – enero 2027).")
    add_bullet(doc, "**Ceremonias:** Sprint Planning, Daily Standup (15 min), Sprint Review, Sprint Retrospective y Backlog Refinement.")
    add_bullet(doc, "**Roles:** Scrum Master (Chiroy Vásquez, Luis David), Product Owner simulado (mismo), Development Team (5 integrantes).")
    add_bullet(doc, "**Definition of Done:** código + pruebas + code review + documentación + despliegue verificado en staging.")
    add_bullet(doc, "**Control de versiones:** Git con ramas `feature/*`, integración en `develop`, producción en `master` protegida, PR obligatorio con al menos 1 aprobación.")
    add_bullet(doc, "**Comunicación:** WhatsApp (coordinación) + GitHub Issues/PRs (código y decisiones técnicas) + Google Meet (ceremonias).")
    add_bullet(doc, "**Calidad:** TDD, ESLint + Prettier, Jest + Supertest, Playwright + axe, pipeline CI de 6 jobs, umbral de cobertura 49/40/48/49 %.")
    add_bullet(doc, "**Honestidad técnica:** lo no implementado se declara `[NO IMPLEMENTADO]` en la documentación, en lugar de afirmarse.")
    add_hr(doc)
